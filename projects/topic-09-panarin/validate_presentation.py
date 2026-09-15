from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
BASE=Path(__file__).resolve().parent
prs=Presentation(BASE/'presentation.pptx')
if len(prs.slides)<10: raise SystemExit(f'Expected >=10 slides, got {len(prs.slides)}')
images=0; violations=[]
for si,slide in enumerate(prs.slides,1):
    for shape in slide.shapes:
        if shape.shape_type==MSO_SHAPE_TYPE.PICTURE: images+=1
        if shape.left<0 or shape.top<0 or shape.left+shape.width>prs.slide_width or shape.top+shape.height>prs.slide_height:
            violations.append((si,shape.name))
if images<4: raise SystemExit(f'Expected >=4 images, got {images}')
if violations: raise SystemExit(f'Out-of-bounds shapes: {violations}')
print(f'OK: {len(prs.slides)} slides, {images} images, no out-of-bounds shapes')
