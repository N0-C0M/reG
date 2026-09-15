from pathlib import Path
import json
import subprocess
import sys
import requests
from docx import Document
from docx.shared import Pt

BASE = Path(__file__).resolve().parent
ASSETS = BASE / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
SOURCES = json.loads((BASE / 'image_sources.json').read_text(encoding='utf-8'))
FILE_MAP = {'portrait':'portrait.jpg','lavra':'lavra.jpg','cathedral':'cathedral.jpg','cathedral_historic':'cathedral_historic.jpg','liturgy':'liturgy.jpg','lavra_liturgy':'lavra_liturgy.jpg'}
headers = {'User-Agent':'Mozilla/5.0 presentation-agent/1.0'}
for key, filename in FILE_MAP.items():
    response = requests.get(SOURCES[key]['url'], headers=headers, timeout=45, allow_redirects=True)
    response.raise_for_status()
    if not response.headers.get('content-type','').startswith('image/'):
        raise RuntimeError(f'{key}: URL did not return an image')
    (ASSETS / filename).write_bytes(response.content)
    print(f'downloaded {key}: {len(response.content)} bytes')
subprocess.run([sys.executable, str(BASE / 'build_presentation.py')], check=True)
md = (BASE / 'speaker_script.md').read_text(encoding='utf-8')
doc = Document(); doc.styles['Normal'].font.name='Arial'; doc.styles['Normal'].font.size=Pt(12)
for raw in md.splitlines():
    line=raw.strip()
    if not line: continue
    if line.startswith('# '): doc.add_heading(line[2:], level=0)
    elif line.startswith('## '): doc.add_heading(line[3:], level=1)
    else: doc.add_paragraph(line)
doc.save(BASE / 'speaker_script.docx')
print('outputs created')
