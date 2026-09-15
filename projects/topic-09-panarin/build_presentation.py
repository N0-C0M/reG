from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from PIL import Image

BASE = Path(__file__).resolve().parent
ASSETS = BASE / 'assets'
OUT = BASE / 'presentation.pptx'
BLUE = RGBColor(20, 65, 145)
ORANGE = RGBColor(226, 126, 33)
DARK = RGBColor(35, 45, 58)
LIGHT = RGBColor(237, 244, 253)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def text(slide, value, x, y, w, h, size=20, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(4)
    tf.margin_top = tf.margin_bottom = Pt(3)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value; r.font.name = 'Arial'; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return sh

def bullets(slide, items, x, y, w, h, size=18):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = '• ' + item; p.space_after = Pt(7)
        r = p.runs[0]; r.font.name = 'Arial'; r.font.size = Pt(size); r.font.color.rgb = DARK
    return sh

def title(slide, value):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.4), Inches(.22), Inches(12.5), Inches(.7))
    s.fill.solid(); s.fill.fore_color.rgb = BLUE; s.line.fill.background()
    text(slide, value, .62, .34, 11.9, .42, 23, True, WHITE)
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.4), Inches(.94), Inches(3.6), Inches(.07))
    s.fill.solid(); s.fill.fore_color.rgb = ORANGE; s.line.fill.background()

def picture(slide, filename, x, y, w, h):
    path = ASSETS / filename
    with Image.open(path) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    target = w / h; ratio = iw / ih
    if ratio > target:
        visible = target / ratio; pic.crop_left = pic.crop_right = (1-visible)/2
    elif ratio < target:
        visible = ratio / target; pic.crop_top = pic.crop_bottom = (1-visible)/2

def content_slide(head, points, image=None, callout=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, head)
    if image:
        bullets(s, points, .72, 1.28, 6.15, 5.55, 17.5)
        picture(s, image, 7.25, 1.35, 5.15, 3.65)
        if callout:
            b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.25), Inches(5.18), Inches(5.15), Inches(.9))
            b.fill.solid(); b.fill.fore_color.rgb = LIGHT; b.line.color.rgb = BLUE
            text(s, callout, 7.45, 5.35, 4.75, .45, 16, True, BLUE, PP_ALIGN.CENTER)
    else:
        bullets(s, points, .85, 1.35, 11.45, 5.55, 18)
    return s

# 1 — title
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(248, 250, 255)
for y, color, h in [(0, BLUE, 1.0), (6.7, ORANGE, .8)]:
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(y), prs.slide_width, Inches(h)); r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background()
text(s, 'Православная цивилизация\nв концепции А. С. Панарина', .8, 1.45, 11.5, 1.2, 28, True, BLUE)
text(s, 'Тема 9 • Российская цивилизация в академическом дискурсе', .82, 3.0, 11.2, .5, 18, True, ORANGE)
text(s, 'Студент: ____________________   Группа: ____________________', .85, 5.75, 11.4, .5, 17, True, WHITE)

content_slide('1. А. С. Панарин: автор и контекст', [
    'Российский философ и политолог, профессор МГУ.',
    'Исследовал глобализацию, цивилизационные процессы и политическую философию.',
    'Ключевая работа для темы — «Православная цивилизация в глобальном мире» (2002).',
    'Православие у него рассматривается шире религии: как культурная память, этика и общественная солидарность.'
], 'portrait.jpg', 'Вопрос группе: что делает общество отдельной цивилизацией?')

content_slide('2. Центральная идея концепции', [
    'Глобализация не должна стирать культурные различия.',
    'Православная традиция может быть альтернативой культу выгоды и силы.',
    'Экономика должна оставаться частью общества, а не единственной мерой ценности.',
    'Главные темы: достоинство человека, солидарность, историческая память и общее благо.'
], 'lavra.jpg')

s = prs.slides.add_slide(prs.slide_layouts[6]); title(s, '3. Реальные фотографии по теме')
for f, x, y in [('cathedral.jpg', .7, 1.25), ('liturgy.jpg', 6.8, 1.25), ('cathedral_historic.jpg', .7, 3.9), ('lavra_liturgy.jpg', 6.8, 3.9)]:
    picture(s, f, x, y, 5.8, 2.25)

content_slide('4. Что такое «православная цивилизация»?', [
    'Не только конфессия и не только набор догматов.',
    'Культурный мир: представления о человеке, добре, справедливости, истории и общине.',
    'Важны память поколений и нравственные ограничения власти и богатства.',
    'Человек понимается как личность, включённая в отношения ответственности перед другими.'
])

content_slide('5. Базовые ценности', [
    'Соборность — единство людей без уничтожения личности.',
    'Служение — противоположность абсолютному эгоизму.',
    'Милосердие — внимание к слабому и беззащитному.',
    'Традиция — память, позволяющая обществу сохранять смысл.',
    'Личная ответственность и нравственные пределы власти.'
], 'cathedral.jpg')

content_slide('6. «Цивилизация бедных»', [
    'Панарин обращает внимание на тех, кто проигрывает от глобального рынка.',
    'Он критикует порядок, в котором сильные приватизируют выгоды, а общество несёт издержки.',
    'Защита слабых становится нравственным критерием общественного устройства.',
    'Это нормативный тезис автора, а не описание всех православных обществ.'
], 'liturgy.jpg')

content_slide('7. Критика глобализации', [
    'Глобализация несёт обмен и открытость, но также риск культурной унификации.',
    'Мобильные элиты и капитал могут выходить из системы национальной ответственности.',
    'Панарин опасается «экономического тоталитаризма», когда всё измеряется выгодой.',
    'Альтернатива — диалог цивилизаций и сохранение множественности моделей развития.'
], 'cathedral_historic.jpg', 'Голосование: глобализация больше объединяет или унифицирует?')

content_slide('8. Россия как цивилизационный субъект', [
    'Россия у Панарина — не «недозападная» страна, а носитель собственного исторического опыта.',
    'Она соединяет разные культурные влияния и сохраняет самостоятельную траекторию.',
    'Модернизация не должна автоматически означать культурную унификацию.',
    'Ключевой вопрос: можно ли развиваться, не теряя историческую идентичность?'
], 'lavra.jpg')

content_slide('9. Религия и публичная этика', [
    'Религиозная традиция может давать обществу язык достоинства, долга и ответственности.',
    'Церковь рассматривается как один из институтов исторической памяти и солидарности.',
    'Современное государство при этом остаётся светским.',
    'Важна граница между культурным влиянием и политическим принуждением.'
], 'lavra_liturgy.jpg')

content_slide('10. Сильные стороны и ограничения', [
    'Сильная сторона: внимание к неравенству и культурному разнообразию.',
    'Сильная сторона: возвращение нравственных вопросов в политическую философию.',
    'Ограничение: риск идеализации традиции.',
    'Ограничение: широкие цивилизационные обобщения могут скрывать внутреннее разнообразие.',
    'Ограничение: трудно напрямую перевести нормативную концепцию в политическую практику.'
])

content_slide('11. Почему тема остаётся актуальной?', [
    'Споры о глобализации, суверенитете и идентичности продолжаются.',
    'Общества всё ещё ищут баланс между открытостью и самобытностью.',
    'Вопрос о социальных издержках глобальной экономики остаётся важным.',
    'Панарин интересен как автор, связывающий культурную идентичность с социальной справедливостью.'
])

content_slide('12. Интерактив: мнение аудитории', [
    'Можно ли сохранить культурную самобытность, не закрываясь от мира?',
    'Должен ли рынок быть высшей мерой общественной жизни?',
    'Нужна ли современному обществу опора на традицию?',
    'Может ли Россия предлагать миру собственную модель развития?'
])

content_slide('13. Мини-кейс', [
    'Представьте глобальную платформу, которая унифицирует образование, медиа и культурный контент.',
    'Назовите один плюс глобального стандарта.',
    'Назовите один риск для локальной культуры и общественной солидарности.',
    'Как бы Панарин оценил такой сценарий?'
])

content_slide('14. Проверка понимания', [
    'Что Панарин противопоставляет «экономическому тоталитаризму»?',
    'Почему для него важна тема защиты слабых?',
    'Как он понимает роль России в глобальном мире?',
    'Назовите две ценности православной цивилизации в его концепции.'
])

content_slide('15. Основные работы', [
    '«Православная цивилизация в глобальном мире» — 2002.',
    '«Искушение глобализмом» — 2000.',
    '«Глобальное политическое прогнозирование» — 2000.',
    '«Россия в циклах мировой истории» — 1999.',
    '«Стратегическая нестабильность XXI века» — 2003.'
])

content_slide('16. Источники', [
    'Панарин А. С. «Православная цивилизация в глобальном мире». Москва, 2002.',
    'Панарин А. С. «Искушение глобализмом». Москва, 2000.',
    'Панарин А. С. «Глобальное политическое прогнозирование». Москва, 2000.',
    'Философский факультет МГУ. Биографическая справка об А. С. Панарине.',
    'Фотографии: Wikimedia Commons, РУВИКИ и официальный сайт ОВЦС Московского Патриархата.'
])

s = prs.slides.add_slide(prs.slide_layouts[6])
for y, color, h in [(0, BLUE, 1.0), (6.7, ORANGE, .8)]:
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(y), prs.slide_width, Inches(h)); r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background()
text(s, 'Спасибо за внимание!', 1, 2.1, 11.3, .7, 30, True, BLUE, PP_ALIGN.CENTER)
text(s, 'Итог: у Панарина православная цивилизация — это культурная память, нравственная ответственность и солидарность в условиях глобализации.', 1.35, 3.25, 10.6, 1.2, 21, False, DARK, PP_ALIGN.CENTER)
text(s, 'Вопросы?', 1, 5.8, 11.3, .4, 21, True, WHITE, PP_ALIGN.CENTER)

prs.save(OUT)
print(f'created {OUT} ({len(prs.slides)} slides)')
